"""Excel workbooks for one area sweep.

Two shapes, chosen by what the sweep actually is:

  * One provider  -> an inventory: a single POI sheet plus a Summary.
  * Two or more   -> a comparison: one sheet per presence bucket ("Only in HERE",
                     "In HERE + GRAB", ...) plus a Summary. Each bucket becomes its
                     own tab so the "only in X" lists are separate documents in the
                     way people actually use them, without producing several files.

CSV stays available for anyone piping the data elsewhere; this is the human-facing
export, with frozen headers, autofilter, sane column widths and the run's provenance
recorded on its own sheet.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import AreaCluster, AreaSweep, PresenceBucket, Provider, ProviderPlace
from .taxonomy import CATEGORY_LABELS, UnifiedCategory

HEADER_FILL = PatternFill("solid", fgColor="17324D")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
BOLD = Font(bold=True, size=10)
TITLE = Font(bold=True, size=12)

#: Human sheet names per bucket. Excel caps a tab name at 31 characters.
BUCKET_SHEET = {
    PresenceBucket.ALL_THREE.value: "In all three",
    PresenceBucket.HERE_GOOGLE.value: "In HERE + Google",
    PresenceBucket.HERE_GRAB.value: "In HERE + Grab",
    PresenceBucket.GOOGLE_GRAB.value: "In Google + Grab",
    PresenceBucket.HERE_ONLY.value: "Only in HERE",
    PresenceBucket.GOOGLE_ONLY.value: "Only in Google",
    PresenceBucket.GRAB_ONLY.value: "Only in Grab",
}

COLUMNS = [
    ("Name", 38), ("Category", 20), ("Address", 52), ("City", 20),
    ("Latitude", 12), ("Longitude", 12), ("Phone", 18), ("Website", 28),
    ("In HERE", 9), ("In Google", 10), ("In Grab", 9),
    ("HERE name", 30), ("Google name", 30), ("Grab name", 30),
    ("Furthest pair (m)", 17), ("Match score", 12),
    ("Provider place ID", 44),
]


def _label(value: str | None) -> str:
    try:
        return CATEGORY_LABELS[UnifiedCategory(value)]
    except (ValueError, KeyError):
        return value or "Unclassified"


def _write_header(sheet) -> None:
    for index, (title, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


def _row(cluster: AreaCluster, places: dict[int, ProviderPlace]) -> list:
    here = places.get(cluster.here_place_id)
    google = places.get(cluster.google_place_id)
    grab = places.get(cluster.grab_place_id)
    # Google and HERE carry contact detail; Grab returns none in this region.
    primary = google or here or grab
    evidence = next(iter((cluster.match_evidence or {}).values()), {})
    return [
        cluster.display_name or (primary.name if primary else ""),
        _label(cluster.unified_category),
        cluster.address or (primary.address if primary else ""),
        (primary.city if primary else "") or "",
        cluster.lat, cluster.lng,
        (primary.phone_e164 if primary else "") or "",
        (primary.website if primary else "") or "",
        "YES" if here else "", "YES" if google else "", "YES" if grab else "",
        here.name if here else "", google.name if google else "", grab.name if grab else "",
        cluster.max_pair_distance_m, evidence.get("score"),
        (primary.provider_place_id if primary else "") or "",
    ]


def _add_sheet(book: Workbook, title: str, clusters, places) -> None:
    sheet = book.create_sheet(title[:31])
    _write_header(sheet)
    ordered = sorted(clusters, key=lambda c: (c.unified_category, (c.display_name or "").lower()))
    for row_index, cluster in enumerate(ordered, start=2):
        for column_index, value in enumerate(_row(cluster, places), start=1):
            sheet.cell(row=row_index, column=column_index, value=value)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(sheet.max_row, 1)}"


def _add_summary(book: Workbook, sweep: AreaSweep, total: int, per_bucket: dict[str, int]) -> None:
    sheet = book.create_sheet("Summary", 0)
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 92
    line = 1

    def put(label: str, value, bold: bool = False) -> None:
        nonlocal line
        sheet.cell(row=line, column=1, value=label).font = TITLE if bold else BOLD
        sheet.cell(row=line, column=2, value=value)
        line += 1

    providers = list(sweep.providers or [])
    comparison = len(set(providers)) > 1
    put(sweep.label or f"Area sweep {sweep.id}", "", bold=True)
    put("Type", "Comparison" if comparison else "Single-source inventory")
    put("Providers", ", ".join(providers))
    put("Area centre", f"{sweep.center_lat}, {sweep.center_lng}")
    put("Radius (m)", sweep.radius_m)
    put("Categories", "All" if not sweep.categories else ", ".join(sweep.categories))
    put("Status", sweep.status.value)
    put("API calls", sweep.api_calls)
    put("Cost (USD)", "unpriced" if sweep.actual_cost_usd is None else sweep.actual_cost_usd)
    put("Distinct places", total)
    if sweep.partial_providers:
        put("PARTIAL COVERAGE", ", ".join(sweep.partial_providers) +
            " stopped at a budget cap; gaps attributed to them are not reliable.")
    line += 1

    put("Records per provider", "", bold=True)
    for provider, count in (sweep.provider_counts or {}).items():
        put(f"  {provider}", count)
    line += 1

    if comparison:
        put("Places by presence", "", bold=True)
        for value, count in per_bucket.items():
            if count:
                put(f"  {BUCKET_SHEET.get(value, value)}", count)
        line += 1

    put("Places by category", "", bold=True)
    matrix = sweep.category_matrix or {}
    for key in sorted(matrix, key=lambda k: -(matrix[k].get("clusters") or 0)):
        put(f"  {_label(key)}", matrix[key].get("clusters", 0))
    line += 1

    put("How to read this", "", bold=True)
    if comparison:
        put("  Gaps", "A place absent from a provider means that provider's own search did "
                      "not return it inside this circle. It does NOT prove the business is "
                      "absent from that provider — that needs a targeted per-record check.")
    else:
        put("  Scope", "One provider was swept, so this is an inventory of that map's "
                       "records. It makes no claim about what other providers hold.")
    if Provider.GRAB.value in providers:
        put("  Grab", "Grab has no radius search, so its coverage is keyword-derived and is "
                      "a floor rather than a ceiling. Grab also returns no phone or website "
                      "in this region.")


def build_workbook(db: Session, sweep: AreaSweep) -> io.BytesIO:
    """Render one sweep to an in-memory .xlsx."""
    clusters = db.scalars(
        select(AreaCluster).where(AreaCluster.sweep_id == sweep.id).order_by(AreaCluster.id)
    ).all()
    place_ids = {
        pid for c in clusters
        for pid in (c.here_place_id, c.google_place_id, c.grab_place_id) if pid
    }
    places = {
        p.id: p for p in
        (db.scalars(select(ProviderPlace).where(ProviderPlace.id.in_(place_ids))).all()
         if place_ids else [])
    }

    book = Workbook()
    book.remove(book.active)  # drop the default sheet; every sheet below is explicit

    by_bucket: dict[str, list] = {}
    for cluster in clusters:
        by_bucket.setdefault(cluster.bucket.value, []).append(cluster)
    counts = {value: len(rows) for value, rows in by_bucket.items()}

    if len(set(sweep.providers or [])) > 1:
        # One tab per bucket, in the order the interface shows them.
        for value in BUCKET_SHEET:
            if by_bucket.get(value):
                _add_sheet(book, BUCKET_SHEET[value], by_bucket[value], places)
        if not book.sheetnames:
            _add_sheet(book, "Places", clusters, places)
    else:
        _add_sheet(book, "POIs", clusters, places)

    _add_summary(book, sweep, len(clusters), counts)

    stream = io.BytesIO()
    book.save(stream)
    stream.seek(0)
    return stream


def workbook_filename(sweep: AreaSweep) -> str:
    providers = "_".join(sweep.providers or ["area"])
    kind = "comparison" if len(set(sweep.providers or [])) > 1 else "inventory"
    return f"GeoDiscover_{providers}_{int(sweep.radius_m)}m_{kind}_{sweep.id}.xlsx"
